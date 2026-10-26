#!/usr/bin/env python3
"""Reproduce WARA Start-WARAReport outputs on Linux (no Office/COM), with a prioritized section first
and an advisory environment/criticality overlay.

Outputs (in --out-dir):
  Executive Summary Presentation - <customer> - <ts>.pptx   (WARA template, slides filled as Start-WARAReport does)
  WARA Assessment Report - <customer> - <ts>.docx           (narrative report; WARA 1.0.6 has no Word output)
  Assessment-Findings-Report-v1-<ts>.xlsx                   (WARA findings workbook, static pivot summaries)
  Impacted Resources and Recommendations Template <ts>.csv  (as Start-WARAReport)
  environment-classification.json, report-manifest.json

Usage:
  python3 generate_reports.py --run-dir <run folder> [--out-dir <run folder>/reports]
      --environment-evidence <environment-evidence.json> [--classification-config <yaml>]
      [--customer NAME] [--workload NAME] [--include-low]
Requires python-pptx, python-docx, openpyxl, matplotlib, PyYAML. WARA templates are read from the runner
cache (WARA_CACHE or ~/.cache/wara-assessment) unless --ppt-template / --findings-template are given.

Deterministic inputs are only read: run.json, derived/prioritized-findings.json, derived/delta.json,
standard/Expert-Analysis-*.xlsx, standard/WARA-File-*.json.
"""
import argparse, copy, csv, datetime as dt, glob, hashlib, html, json, os, re, shutil, sys
from collections import Counter, OrderedDict, defaultdict

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from environment_classification import classify, load_config, overlay  # noqa: E402
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
import docx
from docx.enum.section import WD_ORIENT
from docx.shared import Pt as DPt, Inches as DInches, RGBColor as DRGB

GENERATOR = "wara-assessment generate_reports.py 0.2.0"
AUTOMESSAGE = "AUTOMATICALLY MODIFIED (Please Review)"
BAND_ORDER = ["P1", "P2", "P3", "P4"]
REPORT_BANDS = BAND_ORDER + ["Accepted"]
IMPACT_COLORS = {"High": "C0392B", "Medium": "E67E22", "Low": "7F8C8D"}
LEVEL_ORDER = ["Required", "Recommended", "Optional", "NotRequired"]

# --------------------------------------------------------------------------- inputs

def read_sheet(wb, name, header_row=12):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    hdr = [c.value for c in ws[header_row]]
    out = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(v not in (None, "") for v in r):
            continue
        out.append({h: v for h, v in zip(hdr, r) if h})
    return out


def cnt(x):
    """delta.json stores finding-key lists; older shapes stored counts."""
    return len(x) if isinstance(x, (list, tuple, set)) else (x or 0)


def as_list(x):
    if x is None:
        return []
    return x if isinstance(x, list) else [x]


def html_to_text(s):
    s = re.sub(r"(?i)<br\s*/?>", "\n", s or "")
    s = re.sub(r"(?i)</(p|li|div|h\d)>", "\n", s)
    s = re.sub(r"(?i)<li[^>]*>", "- ", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = html.unescape(s)
    return "\n".join(l.strip() for l in s.splitlines() if l.strip())


OUTAGE_HEADINGS = ["What happened?", "What went wrong and why?", "How did we respond?",
                   "How are we making incidents like this less likely or less impactful?",
                   "How can customers make incidents like this less impactful?",
                   "How can we make our incident communications more useful?"]


def split_pir(text):
    sections, cur = OrderedDict(), None
    for line in text.splitlines():
        h = next((h for h in OUTAGE_HEADINGS if line.lower().startswith(h.lower()[:-1])), None)
        if h:
            cur = h
            sections[cur] = []
        elif cur:
            sections[cur].append(line)
    return {k: "\n".join(v) for k, v in sections.items()}


def load_run(run_dir):
    j = lambda p: json.load(open(os.path.join(run_dir, p), encoding="utf-8-sig"))
    run = j("run.json")
    pf = j("derived/prioritized-findings.json")
    if pf.get("runId") != run.get("runId") or pf.get("scope") != run.get("scope"):
        raise ValueError("Prioritized findings run/scope does not match run.json.")
    for key, entry in run["outputs"].items():
        path, expected = (entry["path"], entry["sha256"]) if isinstance(entry, dict) else (key, entry)
        if sha(os.path.join(run_dir, path)) != expected:
            raise ValueError(f"Recorded input hash mismatch: {path}")
    delta = j("derived/delta.json") if os.path.exists(os.path.join(run_dir, "derived/delta.json")) else {}
    xlsx = [os.path.join(run_dir, run["outputs"]["analyzerWorkbook"]["path"])]
    wjson = [os.path.join(run_dir, run["outputs"]["collectorJson"]["path"])]
    wb = openpyxl.load_workbook(xlsx[0], data_only=True)
    impacted = [r for r in read_sheet(wb, "4.ImpactedResourcesAnalysis") if r.get("Resource Type")]
    inventory = read_sheet(wb, "2.WorkloadInventory")
    platform = [r for r in read_sheet(wb, "5.PlatformIssuesAnalysis") if r.get("Title")]
    tickets = [r for r in read_sheet(wb, "6.SupportRequestsAnalysis") if r.get("Ticket ID")]
    collector = json.load(open(wjson[0], encoding="utf-8")) if wjson else {}
    platform_source = "Expert-Analysis 5.PlatformIssuesAnalysis"
    if not platform and collector.get("outages"):
        platform_source = "collector WARA-File outages (analyzer sheet had no rows)"
        for o in as_list(collector.get("outages")):
            p = o.get("properties", {})
            sec = split_pir(html_to_text(p.get("description", "")))
            svc = sorted({i.get("impactedService") for i in p.get("impact", []) if i.get("impactedService")})
            platform.append({
                "Tracking ID": o.get("name"), "Title": p.get("title"), "Status": p.get("status"),
                "Start Time": p.get("impactStartTime"), "Mitigation Time": p.get("impactMitigationTime"),
                "Impacted Service": ", ".join(svc),
                "What happened": sec.get("What happened?", ""),
                "How can customers make incidents like this less impactful": sec.get("How can customers make incidents like this less impactful?", ""),
            })
    return dict(run=run, pf=pf, delta=delta, impacted=impacted, inventory=inventory, platform=platform,
                platform_source=platform_source, tickets=tickets, collector=collector,
                xlsx=xlsx[0], wjson=wjson[0] if wjson else None)

# --------------------------------------------------------------------------- WARA-equivalent aggregations

def recommendations_table(impacted):
    """Initialize-ExcelRecommendations: one row per Guid (custom rows without Guid kept individually)."""
    groups, out = OrderedDict(), []
    for r in impacted:
        if r.get("Guid"):
            groups.setdefault(r["Guid"], []).append(r)
    for guid, rs in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        f = rs[0]
        out.append({"Impact": f.get("Impact"), "Description": f.get("Recommendation Title"),
                    "Potential Benefit": f.get("Potential Benefit"), "Impacted Resources": len(rs),
                    "Resource Type": f.get("Resource Type"), "Recommendation Control": f.get("Recommendation Control"),
                    "Long Description": f.get("Long Description"), "Category": f.get("Category"),
                    "Learn More Link": f.get("Learn More Link"), "Guid": guid, "Notes": f.get("Notes")})
    for f in impacted:
        if not f.get("Guid"):
            out.append({"Impact": f.get("Impact"), "Description": f.get("Recommendation Title"),
                        "Potential Benefit": f.get("Potential Benefit"), "Impacted Resources": 1,
                        "Resource Type": f.get("Resource Type"), "Recommendation Control": f.get("Recommendation Control"),
                        "Long Description": f.get("Long Description"), "Category": f.get("Category"),
                        "Learn More Link": f.get("Learn More Link"), "Guid": None, "Notes": f.get("Notes")})
    return out


def ci_counter(keys):
    """Case-insensitive Counter (PowerShell Group-Object and Excel pivots ignore case); first-seen spelling is displayed."""
    label, c = {}, Counter()
    for k in keys:
        n = tuple(str(x or "").lower() for x in k) if isinstance(k, tuple) else str(k or "").lower()
        label.setdefault(n, k)
        c[n] += 1
    return Counter({label[n]: v for n, v in c.items()})


def resource_types(impacted, top=10):
    """Slide 12: unique (Resource Type, id) pairs grouped by type, top 10."""
    seen = {}
    for r in impacted:
        seen.setdefault((str(r.get("Resource Type") or "").lower(), str(r.get("id") or "").lower()), r.get("Resource Type"))
    return ci_counter(seen.values()).most_common(top)


def impact_groups(impacted, impact):
    """Slides 23-25: group by (Recommendation Title, Resource Type), sorted by count desc."""
    c = ci_counter((r.get("Recommendation Title"), r.get("Resource Type")) for r in impacted if r.get("Impact") == impact)
    return sorted(c.items(), key=lambda kv: -kv[1])


def pivot(recs, row_key):
    t, label = defaultdict(Counter), {}
    for r in recs:
        k = r.get(row_key) or "(blank)"
        k = label.setdefault(str(k).lower(), k)
        t[k][r.get("Impact") or "(blank)"] += 1
    return t


def dashboard_numbers(impacted):
    u = lambda rs, k: len({str(r.get(k) or "").lower() for r in rs})
    return dict(total=u(impacted, "Guid"),
                high=u([r for r in impacted if r.get("Impact") == "High"], "Guid"),
                medium=u([r for r in impacted if r.get("Impact") == "Medium"], "Guid"),
                low=u([r for r in impacted if r.get("Impact") == "Low"], "Guid"),
                resources=u(impacted, "id"))

# --------------------------------------------------------------------------- PPTX helpers

def shape_by_id(slide, sid):
    return next((s for s in slide.shapes if s.shape_id == sid), None)


def set_para(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run().text = text
    for br in p._p.findall(qn("a:br")):
        p._p.remove(br)


def set_tf(tf, text):
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    set_para(tf.paragraphs[0], text)


def set_tf_paras(tf, items):
    """Fill paragraphs 1..n keeping each paragraph's formatting; extra paragraphs cloned from the last."""
    paras = list(tf.paragraphs)
    while len(paras) < len(items):
        new = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new)
        paras = list(tf.paragraphs)
    for p, t in zip(paras, items):
        set_para(p, t)
    for p in paras[len(items):]:
        p._p.getparent().remove(p._p)


def slide_ids(prs):
    return prs.slides._sldIdLst


def delete_slide(prs, slide):
    lst = slide_ids(prs)
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
            return


def move_slide(prs, slide, index):
    lst = slide_ids(prs)
    el = next(s for s in lst if prs.part.related_part(s.rId) is slide.part)
    lst.remove(el)
    lst.insert(index, el)


def slide_index(prs, slide):
    return next(i for i, s in enumerate(prs.slides) if s.part is slide.part)


def duplicate_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rId, rel in list(src.part.rels.items()):
        if rel.reltype.endswith("/notesSlide") or rel.reltype.endswith("/slideLayout"):
            continue
        rid_map[rId] = (new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external
                        else new.part.rels.get_or_add(rel.reltype, rel.target_part))
    sp_tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        sp_tree.append(copy.deepcopy(el))
    for el in sp_tree.iter():
        for a in (qn("r:embed"), qn("r:id"), qn("r:link")):
            if el.get(a) in rid_map:
                el.set(a, rid_map[el.get(a)])
    bg = src.part._element.cSld.bg
    if bg is not None:
        new.part._element.cSld.insert(0, copy.deepcopy(bg))
    move_slide(prs, new, slide_index(prs, src) + 1)
    return new


def layout(prs, name):
    return next(l for m in prs.slide_masters for l in m.slide_layouts if l.name == name)


def fit_table(gf, rows_used):
    tbl = gf.table._tbl
    for tr in tbl.tr_lst[rows_used:]:
        tbl.remove(tr)
    gf.height = Emu(sum(int(tr.h) for tr in tbl.tr_lst))


def add_table(slide, header, rows, x, y, w, col_w, font=10, header_font=10):
    gf = slide.shapes.add_table(len(rows) + 1, len(header), x, y, w, Emu(int(Inches(0.3)) * (len(rows) + 1)))
    t = gf.table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Emu(int(w * cw))
    for ci, h in enumerate(header):
        c = t.cell(0, ci)
        c.text = str(h)
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold = Pt(header_font), True
    for ri, row in enumerate(rows, 1):
        for ci, v in enumerate(row):
            c = t.cell(ri, ci)
            c.text = "" if v is None else str(v)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font)
    return gf


def add_note(slide, prs, text, y=None, size=10, color="595959", h=None):
    tb = slide.shapes.add_textbox(Inches(0.5), y if y is not None else prs.slide_height - Inches(0.75),
                                  prs.slide_width - Inches(1.0), h or Inches(0.5))
    tb.text_frame.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size, r.font.color.rgb = Pt(size), RGBColor.from_string(color)
    return tb


def new_title_slide(prs, title, after_index):
    s = prs.slides.add_slide(layout(prs, "Title Only"))
    s.shapes.title.text = title
    move_slide(prs, s, after_index)
    return s


def chunks(seq, n):
    seq = list(seq)
    return [seq[i:i + n] for i in range(0, len(seq), n)] or [[]]

# --------------------------------------------------------------------------- PPTX build

def build_pptx(ctx, template, out_path):
    prs = Presentation(template)
    T = {i: s for i, s in enumerate(prs.slides, 1)}  # template slide numbers (1-based) before any change
    imp = ctx["impacted"]
    W, C = ctx["workload"], ctx["customer"]

    # Remove-PPTSlide1 + title on the new first slide (template slide 2, shape 5)
    delete_slide(prs, T[1])
    set_tf(shape_by_id(T[2], 5).text_frame, f"{C} - {W}")

    # Slide 12 (template 13): workload summary
    s = T[13]
    set_tf(shape_by_id(s, 9).text_frame, AUTOMESSAGE)
    sh8 = shape_by_id(s, 8)
    sh8._element.getparent().remove(sh8._element)
    regions = sorted({(r.get("location") or "").lower() for r in ctx["inventory"] if (r.get("location") or "").lower() not in ("", "global")})
    set_tf(shape_by_id(s, 3).text_frame,
           f"During the engagement, the Workload {W} has been reviewed. The solution is hosted in "
           f"{len(regions)} Azure region{'s' if len(regions) != 1 else ''} ({', '.join(regions) or 'n/a'}), "
           f"and runs the following resource types, which include but are not limited to:")
    types = resource_types(imp)
    gf = shape_by_id(s, 6)
    tbl = gf.table._tbl
    total_w = sum(int(c.w) for c in tbl.tblGrid.gridCol_lst)
    for gc in tbl.tblGrid.gridCol_lst[1:]:  # single-column list (WARA widens column 1)
        tbl.tblGrid.remove(gc)
    for tr in tbl.tr_lst:
        for tc in tr.tc_lst[1:]:
            tr.remove(tc)
    tbl.tblGrid.gridCol_lst[0].w = total_w
    proto = copy.deepcopy(tbl.tr_lst[0])
    for i, (t, n) in enumerate(types):
        if i > 0:
            tbl.append(copy.deepcopy(proto))
        cell = gf.table.cell(i, 0)
        set_tf(cell.text_frame, f"{t} ({n})")
    fit_table(gf, max(1, len(types)))

    # Slide 16 (template 17): dashboard numbers + key high impact (Azure Service) titles
    s = T[17]
    set_tf(shape_by_id(s, 41).text_frame, AUTOMESSAGE)
    hi = Counter(r.get("Recommendation Title") for r in imp if r.get("Impact") == "High" and r.get("Category") == "Azure Service")
    set_tf_paras(shape_by_id(s, 9).text_frame, [t for t, _ in hi.most_common(6)] or ["No high impact Azure Service recommendations"])
    d = ctx["dash"]
    for sid, v in ((19, d["total"]), (21, d["high"]), (23, d["medium"]), (24, d["low"]), (30, d["resources"])):
        set_tf(shape_by_id(s, sid).text_frame, str(v))

    # Slide 17 (template 18): charts (Start-WARAReport pastes Excel pivot charts P0/P1; native charts here)
    s = T[18]
    set_tf(shape_by_id(s, 41).text_frame, AUTOMESSAGE)
    impacts = ["High", "Medium", "Low"]
    for (pv, title, x) in ((pivot(ctx["recs"], "Resource Type"), "Recommendations by Impact per ResourceType", Inches(0.4)),
                           (pivot(ctx["recs"], "Recommendation Control"), "Recommendations by Impact per Category", Inches(6.75))):
        if not pv:
            add_note(s, prs, f"{title}: none reported", y=Inches(1.35), size=12)
            continue
        cd = CategoryChartData()
        cats = sorted(pv.keys(), key=lambda k: -sum(pv[k].values()))
        cd.categories = cats
        for im in impacts:
            cd.add_series(im, [pv[c][im] for c in cats])
        ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, Inches(1.35), Inches(6.2), Inches(5.6), cd).chart
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.TOP, False
        ch.category_axis.tick_labels.font.size = Pt(8)
        ch.value_axis.tick_labels.font.size = Pt(8)
        ch.category_axis.reverse_order = True
        for ser, im in zip(ch.plots[0].series, impacts):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = RGBColor.from_string(IMPACT_COLORS[im])

    # Slide 30 (template 31): retirements
    s = T[31]
    ret = [r for r in imp if r.get("Source") == "Azure Service Health - Service Retirements"]
    set_tf(shape_by_id(s, 4).text_frame, AUTOMESSAGE)
    set_tf(shape_by_id(s, 7).text_frame,
           "\n".join(f"{r.get('Retirement TrackingId')} - {r.get('Recommendation Title')}" for r in ret)
           or "No service retirement notifications were reported for the assessed scope.")
    if ret:
        set_tf_paras(shape_by_id(s, 7).text_frame, [f"{r.get('Retirement TrackingId')} - {r.get('Recommendation Title')}" for r in ret[:10]])

    # Slide 29 (template 30): Sev-A support requests (deleted when none, as Start-WARAReport does)
    s = T[30]
    if ctx["tickets"]:
        set_tf(shape_by_id(s, 4).text_frame, AUTOMESSAGE)
        set_tf_paras(shape_by_id(s, 7).text_frame,
                     [x for t in ctx["tickets"][:3] for x in (f"{t.get('Ticket ID')} - {t.get('Title')}", f"Status: {t.get('Status')}", f"Creation Date: {t.get('Creation Date')}")])
    else:
        delete_slide(prs, s)

    # Slide 28 (template 29): recent Microsoft outages, one slide per outage
    s = T[29]
    outs = ctx["platform"]
    set_tf(shape_by_id(s, 4).text_frame, AUTOMESSAGE)
    pages = [s] + [None] * (len(outs) - 1)
    for i in range(1, len(outs)):
        pages[i] = duplicate_slide(prs, pages[i - 1])
    if not outs:
        set_tf_paras(shape_by_id(s, 7).text_frame, ["No Microsoft outages affecting the assessed scope were reported."])
    for page, o in zip(pages, outs):
        tf = shape_by_id(page, 7).text_frame
        paras = list(tf.paragraphs)
        vals = {0: f"{o.get('Tracking ID')} - {o.get('Title')}", 2: o.get("What happened") or "",
                4: o.get("Impacted Service") or "", 6: o.get("How can customers make incidents like this less impactful") or ""}
        for i, v in vals.items():
            if i < len(paras):
                set_para(paras[i], v)
        for p in paras[7:]:
            p._p.getparent().remove(p._p)

    # Slides 23/24/25 (template 24/25/26): High / Medium / Low impact tables, 13 rows per slide
    for tnum, impact in ((24, "High"), (25, "Medium"), (26, "Low")):
        base = T[tnum]
        set_tf(shape_by_id(base, 41).text_frame, AUTOMESSAGE)
        groups = impact_groups(imp, impact)
        pages_rows = chunks(groups, 13)
        pages = [base]
        for _ in pages_rows[1:]:
            pages.append(duplicate_slide(prs, pages[-1]))
        n = 1
        for page, rows in zip(pages, pages_rows):
            gf = shape_by_id(page, 6)
            for ri in range(1, len(gf.table.rows)):
                for ci in range(4):
                    set_tf(gf.table.cell(ri, ci).text_frame, "")
            if not rows:
                set_tf(gf.table.cell(1, 1).text_frame, f"No {impact.lower()} impact recommendations were found.")
                fit_table(gf, 2)
                continue
            for ri, ((title, rtype), count) in enumerate(rows, 1):
                for ci, v in enumerate((n, title, rtype, count)):
                    set_tf(gf.table.cell(ri, ci).text_frame, str(v))
                n += 1
            fit_table(gf, len(rows) + 1)

    # ---- Prioritized section inserted after the title slide (index 1 onwards)
    pos = 1
    sec = duplicate_slide(prs, T[14])  # "Executive Summary" section divider
    set_tf(shape_by_id(sec, 5).text_frame, "Prioritized Reliability Findings")
    set_tf(shape_by_id(sec, 3).text_frame, "Deterministic priority bands (P1-P4) with environment and criticality context")
    move_slide(prs, sec, pos); pos += 1

    run, pf, cls = ctx["run"], ctx["pf"], ctx["cls"]
    sm = pf["summary"]
    scope_cls = ctx["scope_cls"]
    s = new_title_slide(prs, "Assessment Scope and Environment Context", pos); pos += 1
    rows = [
        ("Scope", f"{run['scope']['type']} {run['scope'].get('resourceGroup') or ''} (subscription {run['scope']['subscriptionName']})"),
        ("Run", f"{run['runId']} | status {run['status']} | completed {run.get('completedUtc', '')[:19]}Z"),
        ("Upstream", f"WARA {run['upstream']['waraVersion']} @ {run['upstream']['waraCommit'][:12]}; APRL {run['upstream'].get('aprlRef')} @ {(run['upstream'].get('aprlCommit') or '')[:12]}; {run['upstream']['dependencyMode']} mode"),
        ("Subscription offer", cls["subscription"]["offerInterpretation"]),
        ("Environment classification", f"{ctx['cfg']['tiers'][scope_cls['tier']]['label']} - basis: {scope_cls['basis']} (confidence {scope_cls['confidence']})"),
        ("Assessed against", f"{ctx['cfg']['tiers'][ctx['scope_assessed']]['label']}: {ctx['cfg']['profiles'][ctx['scope_assessed']]}"),
        ("Signals ignored", "; ".join(cls["notes"]) or "none"),
        ("Classification conflicts", "; ".join(sorted({c["conflict"] for c in cls["resources"].values()
                                                       if c.get("conflict")})) or "none"),
    ]
    add_table(s, ["Item", "Value"], rows, Inches(0.5), Inches(1.3), prs.slide_width - Inches(1.0), [0.2, 0.8], font=11, header_font=11)
    add_note(s, prs, "The environment overlay is advisory. It does not change the deterministic P1-P4 bands. Tag resources with "
                     "Environment (Prod/DR/UAT/Test/Dev) and Criticality to replace inference with declared intent.")

    s = new_title_slide(prs, "Prioritized Findings Summary", pos); pos += 1
    lvl = defaultdict(Counter)
    for f in ctx["findings"]:
        lvl[f["priority"]["band"]][f["env"]["requirement"]] += 1
    rows = [(b, sm["bands"].get(b, 0), sum(1 for f in ctx["findings"] if f["priority"]["band"] == b and f["evidence"]["classification"] == "OBSERVED"),
             sum(1 for f in ctx["findings"] if f["priority"]["band"] == b and f["evidence"]["classification"] == "UNKNOWN"),
             *(lvl[b][l] for l in LEVEL_ORDER)) for b in REPORT_BANDS]
    rows.append(("Total", sm["findings"], sm["evidence"].get("OBSERVED", 0), sm["evidence"].get("UNKNOWN", 0),
                 *(sum(lvl[b][l] for b in REPORT_BANDS) for l in LEVEL_ORDER)))
    add_table(s, ["Band", "Findings", "Observed", "Unknown (validate)", "Env: Required", "Env: Recommended", "Env: Optional", "Env: Not required"],
              rows, Inches(0.5), Inches(1.3), prs.slide_width - Inches(1.0), [0.08, 0.1, 0.1, 0.16, 0.14, 0.16, 0.13, 0.13], font=12, header_font=11)
    lc = ", ".join(f"{k} {v}" for k, v in sm.get("lifecycle", {}).items())
    add_note(s, prs, f"Lifecycle vs previous run {sm.get('previousRunId') or 'n/a'}: {lc}; no longer detected {sm.get('resolved', 0)}; "
                     f"not reobserved {sm.get('notReobserved', 0)}.\nEnv columns: requirement level of each finding for its classified "
                     f"environment, from environment-classification.yaml.", y=Inches(4.6), h=Inches(1.0), size=12)

    for band in ("P1", "P2", "Accepted"):
        fs = [f for f in ctx["findings"] if f["priority"]["band"] == band]
        if band == "Accepted" and not fs:
            continue
        for pi, page in enumerate(chunks(fs, 8)):
            s = new_title_slide(prs, f"{band} Findings" + (f" ({pi + 1})" if len(fs) > 8 else ""), pos); pos += 1
            if not page:
                add_note(s, prs, f"No {band} findings.", y=Inches(1.4), size=14)
                continue
            rows = [(f["title"], f["resource"]["name"], f["impact"], f["control"], f["evidence"]["classification"],
                     f["lifecycle"]["status"], f"{f['env']['requirement']} ({f['env']['tier']})", f["env"]["action"]) for f in page]
            add_table(s, ["Recommendation", "Resource", "Impact", "Control", "Evidence", "Lifecycle", "Env requirement", "Env-adjusted action"],
                      rows, Inches(0.4), Inches(1.25), prs.slide_width - Inches(0.8), [0.22, 0.17, 0.06, 0.1, 0.08, 0.09, 0.11, 0.17], font=9, header_font=9)

    unk = [f for f in ctx["findings"] if f["evidence"]["classification"] == "UNKNOWN"]
    s = new_title_slide(prs, "Validation Backlog and Changes", pos); pos += 1
    rows = [(b, sum(1 for f in unk if f["priority"]["band"] == b),
             "; ".join(sorted({f["title"] for f in unk if f["priority"]["band"] == b}))[:300]) for b in BAND_ORDER]
    add_table(s, ["Band", "UNKNOWN items", "Recommendations to validate"], rows, Inches(0.5), Inches(1.3),
              prs.slide_width - Inches(1.0), [0.08, 0.14, 0.78], font=10)
    dl = ctx["delta"]
    add_note(s, prs, f"UNKNOWN means WARA could not evaluate the check automatically; it is a validation task, not a confirmed defect.\n"
                     f"Changes since {dl.get('previousRunId') or 'n/a'}: new {cnt(dl.get('new'))}, persisting {cnt(dl.get('persisting'))}, "
                     f"no longer reported {cnt(dl.get('resolved'))}. Not-reobserved or no-longer-reported findings are not proof of remediation.\n"
                     f"AI review notes (INFERRED, not WARA evidence) are kept in derived/review.md when present.",
             y=Inches(5.0), h=Inches(1.3), size=11)

    # add_slide() names new parts from the slide count, which collides after deletions: renumber all slide parts.
    from pptx.opc.packuri import PackURI
    for i, sl in enumerate(prs.slides, 1):
        sl.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
    prs.save(out_path)
    import zipfile
    names = zipfile.ZipFile(out_path).namelist()
    dup = [n for n, c in Counter(names).items() if c > 1]
    if dup:
        sys.exit(f"PPTX package has duplicate parts: {dup}")
    return len(prs.slides)

# --------------------------------------------------------------------------- DOCX build

def chart_png(pv, title, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    cats = sorted(pv.keys(), key=lambda k: sum(pv[k].values()))
    fig, ax = plt.subplots(figsize=(8, max(2.5, 0.45 * len(cats) + 1)))
    left = [0] * len(cats)
    for im in ("High", "Medium", "Low"):
        vals = [pv[c][im] for c in cats]
        ax.barh(cats, vals, left=left, color="#" + IMPACT_COLORS[im], label=im)
        left = [a + b for a, b in zip(left, vals)]
    ax.set_title(title, fontsize=11)
    ax.legend(loc="lower right", fontsize=8)
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def band_png(sm, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(6, 2.6))
    vals = [sm["bands"].get(b, 0) for b in BAND_ORDER]
    ax.bar(BAND_ORDER, vals, color=["#C0392B", "#E67E22", "#F1C40F", "#7F8C8D"])
    for i, v in enumerate(vals):
        ax.text(i, v + 0.2, str(v), ha="center", fontsize=9)
    ax.set_title("Findings by priority band", fontsize=11)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def dtable(doc, header, rows, widths=None, size=8):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
    for r in t.rows:
        for i, c in enumerate(r.cells):
            if widths:
                c.width = DInches(widths[i])
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = DPt(size)
    doc.add_paragraph()
    return t


def build_docx(ctx, out_path, work):
    run, pf, cls, cfg = ctx["run"], ctx["pf"], ctx["cls"], ctx["cfg"]
    sm, imp = pf["summary"], ctx["impacted"]
    doc = docx.Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, DInches(0.7))
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = DPt(10)

    doc.add_heading("Well-Architected Reliability Assessment", 0)
    doc.add_paragraph(f"{ctx['customer']} - {ctx['workload']}").runs[0].bold = True
    doc.add_paragraph(f"Scope: {run['scope']['type']} {run['scope'].get('resourceGroup') or ''} | subscription {run['scope']['subscriptionName']} ({run['scope']['subscriptionId']})\n"
                      f"Run {run['runId']} | status {run['status']} | generated {ctx['generated']}")
    p = doc.add_paragraph("Part 1 contains the prioritized findings and environment context. Part 2 follows the structure of WARA "
                          "Start-WARAReport (workload summary, recommendations dashboard, impact tables, outages, support requests, "
                          "retirements and next steps). This document was produced by the wara-assessment report generator on "
                          "Linux, not by Start-WARAReport; see Appendix C.")
    p.runs[0].italic = True

    # ---------------- Part 1
    doc.add_heading("Part 1 - Prioritized reliability findings", 1)
    doc.add_heading("1.1 Executive summary", 2)
    scope_cls, assessed = ctx["scope_cls"], ctx["scope_assessed"]
    req_counts = Counter(f["env"]["requirement"] for f in ctx["findings"] if f["priority"]["band"] in ("P1", "P2"))
    doc.add_paragraph(
        f"The assessment produced {sm['findings']} findings: P1 {sm['bands'].get('P1', 0)}, P2 {sm['bands'].get('P2', 0)}, "
        f"P3 {sm['bands'].get('P3', 0)}, P4 {sm['bands'].get('P4', 0)}. {sm['evidence'].get('OBSERVED', 0)} are observed in configuration "
        f"and {sm['evidence'].get('UNKNOWN', 0)} need manual validation. Compared with run {sm.get('previousRunId') or 'n/a'}: "
        f"{cnt(ctx['delta'].get('new'))} new, {cnt(ctx['delta'].get('persisting'))} persisting, {cnt(ctx['delta'].get('resolved'))} no longer reported.")
    doc.add_paragraph(
        f"Environment context: the scope is classified as {cfg['tiers'][scope_cls['tier']]['label']} "
        f"(basis: {scope_cls['basis']}; confidence {scope_cls['confidence']}) and assessed against "
        f"{cfg['tiers'][assessed]['label']} resiliency targets. Under that profile, the {sm['bands'].get('P1', 0) + sm['bands'].get('P2', 0)} "
        f"P1/P2 findings break down as: " + ", ".join(f"{l} {req_counts.get(l, 0)}" for l in LEVEL_ORDER) + ".")
    if scope_cls["confidence"] in ("Low", "None"):
        doc.add_paragraph("The classification is inferred from naming tags only. Before accepting risk on Production-grade "
                          "findings, confirm the environment with the workload owner and add explicit Environment and Criticality tags.",
                          style="Intense Quote")
    band_png(sm, os.path.join(work, "bands.png"))
    doc.add_picture(os.path.join(work, "bands.png"), width=DInches(5.5))

    doc.add_heading("1.2 Environment and criticality context", 2)
    sub = cls["subscription"]
    rows = [("Subscription offer (quotaId)", sub["quotaId"], sub["offerInterpretation"])]
    rows += [(f"Subscription tag {k}", v, "") for k, v in sub["tags"].items()]
    for g, c in cls["resourceGroups"].items():
        rows.append((f"Resource group {g}", c["basis"], f"{cfg['tiers'][c['tier']]['label']} (confidence {c['confidence']})"))
    for n in cls["notes"]:
        rows.append(("Ignored signal", n, ""))
    for c in [c for c in cls["resources"].values() if c.get("conflict")]:
        rows.append(("Conflict", c["name"], c["conflict"]))
    dtable(doc, ["Signal", "Value", "Interpretation"], rows, [2.5, 3.5, 3.5], size=9)
    doc.add_paragraph("Classification precedence: explicit resource tag > explicit resource-group tag > explicit subscription tag > "
                      "non-production subscription offer (MSDN / Visual Studio Dev/Test, Internal_*, per config) > naming-tag heuristics (resource, resource group, subscription) "
                      "> unclassified, which is assessed as Production (primary).")
    doc.add_paragraph("Per-resource classification:")
    dtable(doc, ["Resource", "Type", "Tier", "Basis", "Confidence"],
           [(r["name"], r["type"], r["tier"], r["basis"], r["confidence"]) for r in sorted(cls["resources"].values(), key=lambda r: r["name"])],
           [2.8, 2.6, 1.0, 2.6, 0.8], size=8)
    doc.add_paragraph("Resiliency requirements by environment (requirement level per WARA control):")
    controls = ["HighAvailability", "DisasterRecovery", "Scalability", "MonitoringAndAlerting", "ServiceUpgradeAndRetirement", "Security", "OtherBestPractices"]
    tiers = [t for t in ("PROD-PRIMARY", "PROD-DR", "UAT", "TEST", "DEV")]
    dtable(doc, ["Control"] + tiers, [[c] + [cfg["requirements"][t].get(c, cfg["requirements"][t]["default"]) for t in tiers] for c in controls],
           [2.4] + [1.4] * len(tiers), size=8)
    dtable(doc, ["Tier", "Resiliency target"], [(cfg["tiers"][t]["label"], cfg["profiles"][t]) for t in tiers + ["UNKNOWN"]], [2.2, 7.4], size=8)

    doc.add_heading("1.3 Prioritized findings", 2)
    for band in REPORT_BANDS:
        fs = [f for f in ctx["findings"] if f["priority"]["band"] == band]
        doc.add_heading(f"{band} ({len(fs)})", 3)
        if not fs:
            doc.add_paragraph("None.")
            continue
        dtable(doc, ["Recommendation", "Resource", "Impact", "Control", "Evidence", "Lifecycle", "Env requirement", "Env-adjusted action"],
               [(f["title"], f["resource"]["name"], f["impact"], f["control"], f["evidence"]["classification"], f["lifecycle"]["status"],
                 f"{f['env']['requirement']} ({f['env']['tier']})", f["env"]["action"]) for f in fs],
               [2.4, 1.9, 0.6, 1.1, 0.8, 0.9, 1.1, 1.6], size=8)

    doc.add_heading("1.4 Validation backlog", 2)
    unk = [f for f in ctx["findings"] if f["evidence"]["classification"] == "UNKNOWN"]
    doc.add_paragraph(f"{len(unk)} findings could not be evaluated automatically (UNKNOWN). They are validation tasks, not confirmed "
                      "defects. AI review notes (INFERRED, not WARA evidence), when present, are in derived/review.md of the run.")
    dtable(doc, ["Band", "Recommendation", "Resource", "Env requirement"],
           [(f["priority"]["band"], f["title"], f["resource"]["name"], f["env"]["requirement"]) for f in unk], [0.6, 4.4, 3.0, 1.4], size=8)

    doc.add_heading("1.5 Changes since the previous run", 2)
    d = ctx["delta"]
    doc.add_paragraph(f"Previous run: {d.get('previousRunId') or 'n/a'}. New {cnt(d.get('new'))}, persisting {cnt(d.get('persisting'))}, "
                      f"no longer reported {cnt(d.get('resolved'))}, not reobserved {sm.get('notReobserved', 0)}. Findings that are no "
                      "longer reported or not reobserved are not proof of remediation or deletion.")

    # ---------------- Part 2 (Start-WARAReport structure)
    doc.add_page_break()
    doc.add_heading("Part 2 - Well-Architected Reliability Assessment (WARA report structure)", 1)
    doc.add_heading("2.1 Workload summary", 2)
    regions = sorted({(r.get("location") or "").lower() for r in ctx["inventory"] if (r.get("location") or "").lower() not in ("", "global")})
    doc.add_paragraph(f"During the engagement, the workload {ctx['workload']} has been reviewed. The solution is hosted in "
                      f"{len(regions)} Azure region(s) ({', '.join(regions)}); the inventory contains {len(ctx['inventory'])} resources. "
                      "Resource types with impacted resources:")
    types = resource_types(imp)
    dtable(doc, ["Resource type", "Impacted resources"], types, [5.0, 1.5], size=9)

    doc.add_heading("2.2 Recommendations dashboard", 2)
    dsh = ctx["dash"]
    dtable(doc, ["Total recommendations", "High impact", "Medium impact", "Low impact", "Impacted resources"],
           [(dsh["total"], dsh["high"], dsh["medium"], dsh["low"], dsh["resources"])], [1.9] * 5, size=10)
    hi = Counter(r.get("Recommendation Title") for r in imp if r.get("Impact") == "High" and r.get("Category") == "Azure Service")
    doc.add_paragraph("Key high impact issues (Azure resources):")
    for t, _ in hi.most_common(6):
        doc.add_paragraph(t, style="List Bullet")
    for key, title, fn in (("Resource Type", "Recommendations by Impact per ResourceType", "p0.png"),
                           ("Recommendation Control", "Recommendations by Impact per Category", "p1.png")):
        chart_png(pivot(ctx["recs"], key), title, os.path.join(work, fn))
        doc.add_picture(os.path.join(work, fn), width=DInches(7.5))

    doc.add_heading("2.3 Health and risk recommendations", 2)
    for impact in ("High", "Medium", "Low"):
        g = impact_groups(imp, impact)
        doc.add_heading(f"{impact} impact issues ({len(g)})", 3)
        dtable(doc, ["#", "Recommendation", "Service", "Impacted resources"],
               [(i, t, rt, c) for i, ((t, rt), c) in enumerate(g, 1)], [0.4, 5.4, 2.8, 1.2], size=8)

    doc.add_heading("2.4 Recommendation details", 2)
    for r in ctx["recs"]:
        doc.add_heading(f"[{r['Impact']}] {r['Description']}", 4)
        doc.add_paragraph(f"Resource type: {r['Resource Type']} | Control: {r['Recommendation Control']} | "
                          f"Impacted resources: {r['Impacted Resources']} | Category: {r['Category']}")
        if r.get("Long Description"):
            doc.add_paragraph(str(r["Long Description"]))
        if r.get("Potential Benefit"):
            doc.add_paragraph(f"Potential benefit: {r['Potential Benefit']}")
        if r.get("Learn More Link"):
            doc.add_paragraph(f"Learn more: {r['Learn More Link']}")

    doc.add_heading("2.5 Design, platform and support", 2)
    doc.add_heading("Recent Microsoft outages", 3)
    if not ctx["platform"]:
        doc.add_paragraph("No Microsoft outages affecting the assessed scope were reported.")
    for o in ctx["platform"]:
        doc.add_paragraph(f"{o.get('Tracking ID')} - {o.get('Title')}").runs[0].bold = True
        doc.add_paragraph(f"Status: {o.get('Status') or ''} | Impacted service: {o.get('Impacted Service') or ''} | "
                          f"Start: {o.get('Start Time') or ''} | Mitigated: {o.get('Mitigation Time') or ''}")
        if o.get("What happened"):
            doc.add_paragraph("What happened: " + o["What happened"])
        if o.get("How can customers make incidents like this less impactful"):
            doc.add_paragraph("How can customers make incidents like this less impactful: " + o["How can customers make incidents like this less impactful"])
    doc.add_paragraph(f"Source: {ctx['platform_source']}.").runs[0].italic = True
    doc.add_heading("Sev-A support requests", 3)
    if ctx["tickets"]:
        dtable(doc, ["Ticket ID", "Title", "Status", "Creation date"],
               [(t.get("Ticket ID"), t.get("Title"), t.get("Status"), t.get("Creation Date")) for t in ctx["tickets"]], [1.5, 5.0, 1.2, 1.8])
    else:
        doc.add_paragraph("No support requests were reported for the assessed scope.")
    doc.add_heading("Service retirement notifications", 3)
    ret = [r for r in imp if r.get("Source") == "Azure Service Health - Service Retirements"]
    if ret:
        dtable(doc, ["Tracking ID", "Recommendation", "Resource"], [(r.get("Retirement TrackingId"), r.get("Recommendation Title"), r.get("name")) for r in ret], [1.5, 5.0, 3.0])
    else:
        doc.add_paragraph("No service retirement notifications were reported for the assessed scope.")

    doc.add_heading("2.6 Next steps", 2)
    for t in ("Remediation planning: address Required P1/P2 findings first; record documented risk acceptances for Optional and Not-required items in the prioritization config.",
              "Validate UNKNOWN findings and close them or convert them to remediation items.",
              "Tag resources with Environment and Criticality so future runs classify from declared intent rather than inference.",
              "Triage the Azure Advisor reliability recommendations (accept/reject) so tracking continues in Advisor.",
              "Adopt Azure Monitor Baseline Alerts (aka.ms/amba) for service health and resource health alerting.",
              "Re-run the monthly assessment and review deltas."):
        doc.add_paragraph(t, style="List Number")

    # ---------------- Appendices
    doc.add_page_break()
    doc.add_heading("Appendix A - Impacted resources", 1)
    dtable(doc, ["Resource", "Type", "Recommendation", "Impact", "Control", "Source"],
           [(r.get("name"), r.get("Resource Type"), r.get("Recommendation Title"), r.get("Impact"), r.get("Recommendation Control"), r.get("Source")) for r in imp],
           [2.2, 2.0, 3.2, 0.7, 1.1, 0.8], size=7)
    doc.add_heading("Appendix B - Workload inventory", 1)
    dtable(doc, ["Name", "Type", "Location", "SKU", "Zones"],
           [(r.get("name"), r.get("type"), r.get("location"), str(r.get("sku") or "")[:60], r.get("zones")) for r in ctx["inventory"]],
           [2.6, 3.0, 1.0, 2.2, 0.8], size=7)
    doc.add_heading("Appendix C - Method, provenance and deviations", 1)
    up = run["upstream"]
    for t in (f"Inputs: run.json, derived/prioritized-findings.json, derived/delta.json, {os.path.basename(ctx['xlsx'])}, "
              f"{os.path.basename(ctx['wjson'] or '')} (read-only; sha256 values in report-manifest.json).",
              f"Prioritization: {pf.get('generatedBy', '')} with config sha256 {str(pf.get('config', {}).get('sha256', ''))[:12]}. "
              "Score = impact x evidence x control x criticality; bands are retained from the run's prioritization policy.",
              f"WARA {up['waraVersion']} @ {up['waraCommit'][:12]}; APRL {up.get('aprlRef')} @ {(up.get('aprlCommit') or '')[:12]}; {up['dependencyMode']} mode.",
              *[f"Run deviation: {x}" for x in up.get("deviations", [])],
              "Report generation: Start-WARAReport requires Windows with PowerPoint and Excel (COM). This report reproduces its outputs with "
              "python-pptx, python-docx and openpyxl using WARA's own PowerPoint and Assessment Findings templates. Excel pivot "
              "tables/charts are replaced by static summaries and native charts. WARA 1.0.6 produces no Word document; the Word "
              "report is an addition that follows the presentation's structure.",
              f"Environment classification: evidence collected {cls.get('evidenceCollectedUtc')} by {cls.get('evidenceCollectedBy')}; "
              "rules in environment-classification.yaml. Advisory only; deterministic bands unchanged.",
              f"Platform issues source: {ctx['platform_source']}."):
        doc.add_paragraph(t, style="List Bullet")
    doc.save(out_path)

# --------------------------------------------------------------------------- Assessment Findings workbook + CSV

IR_COLS = ["Impacted?", "Resource Type", "subscriptionId", "resourceGroup", "location", "name", "id", "custom1", "custom2", "custom3",
           "custom4", "custom5", "Recommendation Title", "Impact", "Recommendation Control", "Potential Benefit", "Learn More Link",
           "Long Description", "Guid", "Category", "Source", "WAF Pillar", "Platform Issue TrackingId", "Retirement TrackingId",
           "Support Request Number", "Notes", "checkName"]
REC_COLS = ["Impact", "Description", "Potential Benefit", "Impacted Resources", "Resource Type", "Recommendation Control",
            "Long Description", "Category", "Learn More Link", "Guid", "Notes"]
INV_COLS = ["id", "name", "type", "tenantId", "kind", "location", "resourceGroup", "subscriptionId", "managedBy", "sku", "plan", "zones"]


def xl(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v)
    return v


def write_block(ws, start_row, cols, rows):
    for ci, c in enumerate(cols, 1):
        ws.cell(start_row, ci, c)
    for ri, r in enumerate(rows, start_row + 1):
        for ci, c in enumerate(cols, 1):
            ws.cell(ri, ci, xl(r.get(c)))


def build_xlsx(ctx, template, out_path):
    wb = openpyxl.load_workbook(template)
    get = lambda n: wb[n] if n in wb.sheetnames else wb.create_sheet(n)
    write_block(get("3.ImpactedResources"), 12, IR_COLS, [dict(r, **{"Impacted?": "Yes"}) for r in ctx["impacted"]])
    write_block(get("2.Recommendations"), 11, REC_COLS, ctx["recs"])
    write_block(get("6.WorkloadInventory"), 12, INV_COLS, ctx["inventory"])
    ws = get("7.PivotTable")
    blocks = []
    for col, key in ((1, "Resource Type"), (8, "Recommendation Control")):
        pv = pivot(ctx["recs"], key)
        ws.cell(3, col, f"Count of recommendations - {key} x Impact (static; Start-WARAReport uses a pivot table)")
        hdr = [key, "High", "Medium", "Low", "Grand Total"]
        for i, h in enumerate(hdr):
            ws.cell(4, col + i, h)
        cats = sorted(pv)
        for ri, c in enumerate(cats, 5):
            vals = [c, pv[c]["High"], pv[c]["Medium"], pv[c]["Low"], sum(pv[c].values())]
            for i, v in enumerate(vals):
                ws.cell(ri, col + i, v)
        blocks.append((col, len(cats), key))
    for col, key, agg in ((15, "Impacted Resources", "Sum"), (19, "Guid", "Count")):
        ws.cell(3, col, f"{agg} of {key} by Impact (static)")
        ws.cell(4, col, "Impact"); ws.cell(4, col + 1, f"{agg} of {key}")
        for ri, im in enumerate(("High", "Medium", "Low"), 5):
            rs = [r for r in ctx["recs"] if r["Impact"] == im]
            ws.cell(ri, col, im)
            ws.cell(ri, col + 1, sum(r["Impacted Resources"] for r in rs) if agg == "Sum" else sum(1 for r in rs if r.get("Guid")))
    dash = get("1.Dashboard")
    for (col, n, key), anchor, title in zip(blocks, ("B18", "J18"), ("Recommendations by Impact per ResourceType", "Recommendations by Impact per Category")):
        if n == 0:
            continue
        ch = BarChart()
        ch.type, ch.grouping, ch.title = "bar", "clustered", title
        ch.add_data(Reference(ws, min_col=col + 1, max_col=col + 3, min_row=4, max_row=4 + n), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=col, min_row=5, max_row=4 + n))
        ch.height, ch.width = 14, 16
        dash.add_chart(ch, anchor)
    wb.save(out_path)


def build_csv(ctx, out_path, include_low=False):
    imps = ("High", "Medium", "Low") if include_low else ("High", "Medium")
    with open(out_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Recommendation Guid", "Recommendation Title", "Description", "Priority", "Customer-facing annotation",
                    "Internal-facing notes", "Potential Benefit", "Resource Type", "Resource ID"])
        for r in ctx["impacted"]:
            if r.get("Impact") in imps and r.get("ValidationCategory") != "Retirements":
                w.writerow([r.get("Guid"), r.get("Recommendation Title"), r.get("Long Description"), r.get("Impact"), "",
                            r.get("Notes"), r.get("Potential Benefit"), r.get("Resource Type"), r.get("id")])

# --------------------------------------------------------------------------- main

def find_template(name, explicit, commit):
    if explicit:
        return explicit
    cache = os.environ.get("WARA_CACHE") or os.path.expanduser("~/.cache/wara-assessment")
    path = os.path.join(cache, "src", f"wara-{commit[:12]}", "src", "modules", "wara", "reports", name)
    if not os.path.isfile(path):
        sys.exit(f"Template '{name}' not found under {cache}; run Invoke-WaraRun.ps1 once (restricted mode caches WARA source) or pass it explicitly.")
    return path


def sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run-dir", required=True)
    ap.add_argument("--out-dir", help="default: <run-dir>/reports")
    ap.add_argument("--customer", default="[Customer Name]")
    ap.add_argument("--workload", default="[Workload Name]")
    ap.add_argument("--classification-config", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "config", "environment-classification.yaml"))
    ap.add_argument("--environment-evidence", required=True)
    ap.add_argument("--ppt-template")
    ap.add_argument("--findings-template")
    ap.add_argument("--include-low", action="store_true")
    a = ap.parse_args()

    a.out_dir = a.out_dir or os.path.join(a.run_dir, "reports")
    if os.path.exists(os.path.join(a.out_dir, "report-manifest.json")):
        raise ValueError("Reports already published here; choose a new --out-dir to preserve prior reports.")
    os.makedirs(a.out_dir, exist_ok=True)
    work = os.path.join(a.out_dir, ".work")
    os.makedirs(work, exist_ok=True)
    cfg = load_config(a.classification_config)
    ev = json.load(open(a.environment_evidence, encoding="utf-8-sig"))
    ctx = load_run(a.run_dir)
    scope = ctx["run"]["scope"]
    ev_scope = ev.get("scope") or {}
    if (str(ev_scope.get("subscriptionId") or "").lower() != str(scope["subscriptionId"]).lower()
            or str(ev_scope.get("resourceGroup") or "").lower() != str(scope.get("resourceGroup") or "").lower()):
        raise ValueError("Environment evidence scope does not match run scope.")
    ctx.update(customer=a.customer, workload=a.workload, cfg=cfg,
               generated=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    ctx["recs"] = recommendations_table(ctx["impacted"])
    ctx["dash"] = dashboard_numbers(ctx["impacted"])
    ctx["cls"] = cls = classify(cfg, ev)
    rg = ctx["run"]["scope"].get("resourceGroup") or ""
    ctx["scope_cls"] = (cls["resourceGroups"].get(rg.lower(), dict(tier="UNKNOWN", basis="resource group not in evidence", confidence="None"))
                        if rg else cls["subscription"])
    ctx["scope_assessed"] = cfg["unknownAssessedAs"] if ctx["scope_cls"]["tier"] == "UNKNOWN" else ctx["scope_cls"]["tier"]
    fs = sorted(ctx["pf"]["findings"], key=lambda f: (REPORT_BANDS.index(f["priority"]["band"]), -f["priority"]["score"], f["title"]))
    for f in fs:
        f["env"] = overlay(cfg, cls, f, rg)
    ctx["findings"] = fs

    ts = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d-%H-%M")
    safe = re.sub(r"[^A-Za-z0-9 ._-]", "_", a.customer)
    out = {
        "pptx": os.path.join(a.out_dir, f"Executive Summary Presentation - {safe} - {ts}.pptx"),
        "docx": os.path.join(a.out_dir, f"WARA Assessment Report - {safe} - {ts}.docx"),
        "xlsx": os.path.join(a.out_dir, f"Assessment-Findings-Report-v1-{ts}.xlsx"),
        "csv": os.path.join(a.out_dir, f"Impacted Resources and Recommendations Template {ts}.csv"),
        "classification": os.path.join(a.out_dir, "environment-classification.json"),
    }
    deps_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "config", "dependencies.json")
    deps = json.load(open(deps_path, encoding="utf-8"))
    template_commit = deps["wara"]["gitCommit"]
    ppt_t = find_template("Mandatory - Executive Summary presentation - Template.pptx", a.ppt_template, template_commit)
    xl_t = find_template("Assessment-Findings-Report-v1.xlsx", a.findings_template, template_commit)
    slides = build_pptx(ctx, ppt_t, out["pptx"])
    build_docx(ctx, out["docx"], work)
    build_xlsx(ctx, xl_t, out["xlsx"])
    build_csv(ctx, out["csv"], a.include_low)
    notice = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "config", "wara-report-NOTICE.txt")
    out["notice"] = os.path.join(a.out_dir, "wara-report-NOTICE.txt")
    shutil.copyfile(notice, out["notice"])
    json.dump(dict(cls, scope=dict(resourceGroup=rg, **ctx["scope_cls"], assessedAs=ctx["scope_assessed"]),
                   findings=[dict(findingKey=f["findingKey"], band=f["priority"]["band"], title=f["title"], resource=f["resource"]["name"],
                                  resourceId=f["resource"].get("id"), control=f["control"], **f["env"]) for f in fs]),
              open(out["classification"], "w"), indent=2)
    inputs = [os.path.join(a.run_dir, p) for p in ("run.json", "derived/prioritized-findings.json", "derived/delta.json")] + [ctx["xlsx"]] + ([ctx["wjson"]] if ctx["wjson"] else [])
    manifest = dict(generator=GENERATOR, generatedUtc=ctx["generated"], runId=ctx["run"]["runId"], customer=a.customer, workload=a.workload,
                    pptSlides=slides, platformIssuesSource=ctx["platform_source"],
                    templateSource=dict(repository=deps["wara"]["gitRepository"], pinnedCommit=template_commit,
                                        explicitOverride=bool(a.ppt_template or a.findings_template)),
                    templates={os.path.basename(ppt_t): sha(ppt_t), os.path.basename(xl_t): sha(xl_t)},
                    inputs={os.path.relpath(p, a.run_dir): sha(p) for p in inputs},
                    classificationConfig=sha(a.classification_config), environmentEvidence=sha(a.environment_evidence),
                    outputs={os.path.basename(p): sha(p) for p in out.values()})
    json.dump(manifest, open(os.path.join(a.out_dir, "report-manifest.json"), "w"), indent=2)
    shutil.rmtree(work)
    print(json.dumps(dict(slides=slides, outDir=a.out_dir, outputs=list(manifest["outputs"])), indent=1))


if __name__ == "__main__":
    main()
